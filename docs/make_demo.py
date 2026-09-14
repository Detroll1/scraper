"""Генератор честного демо-пакета в docs/.

Делает реальный запуск CLI на демо-источнике и сохраняет:
  docs/demo_terminal.txt — снимок вывода терминала (реальный stdout, не рисованный);
  docs/demo_table.txt    — текстовая выгрузка получившегося xlsx через openpyxl;
  docs/demo_result.xlsx / docs/demo_result.json — сами результаты.

Запуск из корня проекта:  python docs/make_demo.py
"""

from __future__ import annotations

import os
import subprocess
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DOCS = Path(__file__).resolve().parent

COMMAND = [
    sys.executable, "run.py", "--source", "demo",
    "--out", "docs/demo_result.xlsx", "--limit", "8", "--json", "docs/demo_result.json",
]
COMMAND_LINE = " ".join(COMMAND[1:])
STAMP = datetime.now().strftime("%Y-%m-%d %H:%M")


def render_sheet(sheet: openpyxl.worksheet.worksheet.Worksheet) -> str:
    rows = [
        ["" if cell.value is None else str(cell.value) for cell in row]
        for row in sheet.iter_rows()
    ]
    widths = [
        max(len(row[column]) for row in rows) for column in range(len(rows[0]))
    ]
    lines: list[str] = []
    for index, row in enumerate(rows):
        lines.append(
            " | ".join(cell.ljust(widths[col]) for col, cell in enumerate(row)).rstrip()
        )
        if index == 0:
            lines.append("-+-".join("-" * width for width in widths))
    return "\n".join(lines)


def main() -> int:
    child_env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    process = subprocess.run(
        COMMAND, cwd=ROOT, capture_output=True, text=True,
        encoding="utf-8", env=child_env,
    )

    terminal_lines = [
        "ДЕМО-СНИМОК ТЕРМИНАЛА",
        f"Сгенерирован скриптом docs/make_demo.py {STAMP}. Это реальный вывод программы, не рисованный.",
        f"Воспроизведение: python {COMMAND_LINE}   (из корня проекта)",
        f"Код возврата: {process.returncode}",
        "",
        f"> python {COMMAND_LINE}",
    ]
    terminal_lines += (process.stdout or "").rstrip().splitlines()
    if process.stderr and process.stderr.strip():
        terminal_lines += process.stderr.rstrip().splitlines()
    (DOCS / "demo_terminal.txt").write_text("\n".join(terminal_lines) + "\n", encoding="utf-8")

    workbook = openpyxl.load_workbook(DOCS / "demo_result.xlsx")
    table_lines = [
        "СОДЕРЖИМОЕ ДЕМО-ФАЙЛА docs/demo_result.xlsx",
        f"Выгружено скриптом docs/make_demo.py {STAMP} из реально сгенерированного файла (openpyxl).",
        "",
    ]
    for name in workbook.sheetnames:
        table_lines += [f"Лист «{name}»:", render_sheet(workbook[name]), ""]
    (DOCS / "demo_table.txt").write_text("\n".join(table_lines) + "\n", encoding="utf-8")

    print("Демо-файлы обновлены:")
    print("  docs/demo_terminal.txt — снимок вывода терминала")
    print("  docs/demo_table.txt    — содержимое xlsx в тексте")
    print("  docs/demo_result.xlsx  — сам Excel")
    print("  docs/demo_result.json  — JSON")
    return process.returncode


if __name__ == "__main__":
    sys.exit(main())

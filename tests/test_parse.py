"""Юнит-тесты ядра парсера: разбор записи, цены, дубликаты, запись Excel.

Запуск из папки проекта:
    python -m unittest discover -s tests -v
"""

from __future__ import annotations

import sys
import unittest
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT))

import scraper_core as core  # noqa: E402

FIXTURE = PROJECT_ROOT / "tests" / "fixtures" / "sample.html"
BASE_URL = "https://example.com/"


class ParseRecordsTest(unittest.TestCase):
    """Разбор фикстуры каталога (та же структура, что на books.toscrape.com)."""

    @classmethod
    def setUpClass(cls):
        cls.records = core.parse_records(FIXTURE.read_text(encoding="utf-8"), BASE_URL)

    def test_count_and_dedupe(self):
        # В фикстуре 4 карточки, но «Sea of Stone» повторяется — остаются 3 записи.
        self.assertEqual(len(self.records), 3)

    def test_first_record_fields(self):
        record = self.records[0]
        self.assertEqual(record.title, "Sea of Stone")
        self.assertEqual(record.link, BASE_URL + "catalog/sea-of-stone_995/index.html")
        self.assertEqual(record.price, 51.77)
        self.assertEqual(record.meta, "Рейтинг: Three; In stock (12)")

    def test_meta_without_rating(self):
        third = self.records[2]
        self.assertEqual(third.title, "Winter Tales")
        self.assertEqual(third.meta, "Out of stock")

    def test_titles(self):
        self.assertEqual(
            [record.title for record in self.records],
            ["Sea of Stone", "Silent Ridge", "Winter Tales"],
        )


class ParsePriceTest(unittest.TestCase):
    def test_pound(self):
        self.assertEqual(core.parse_price("£51.77"), 51.77)

    def test_ruble_comma(self):
        self.assertEqual(core.parse_price("499,00 ₽"), 499.0)

    def test_word_currency(self):
        self.assertEqual(core.parse_price("299 руб."), 299.0)

    def test_no_price(self):
        self.assertIsNone(core.parse_price("Нет в наличии"))

    def test_empty(self):
        self.assertIsNone(core.parse_price(None))
        self.assertIsNone(core.parse_price(""))


class GenericRecordsTest(unittest.TestCase):
    def test_generic_fallback(self):
        html = "<html><body><ol><li><a href='/ebooks/84'>Frankenstein</a></li></ol></body></html>"
        records = core.parse_records(html, "https://www.gutenberg.org")
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].title, "Frankenstein")
        self.assertEqual(records[0].link, "https://www.gutenberg.org/ebooks/84")


class WorkbookTest(unittest.TestCase):
    def test_workbook_roundtrip(self):
        import openpyxl

        records = core.parse_records(FIXTURE.read_text(encoding="utf-8"), BASE_URL)
        out_path = PROJECT_ROOT / "tests" / "output_test.xlsx"
        try:
            core.build_workbook(records, out_path, source_label="тест-источник")

            book = openpyxl.load_workbook(out_path)
            self.assertEqual(book.sheetnames, ["Данные", "Сводка"])

            sheet = book["Данные"]
            self.assertEqual(
                [cell.value for cell in sheet[1]],
                ["№", "Наименование", "Ссылка", "Цена", "Мета"],
            )
            self.assertEqual(sheet.max_row, len(records) + 1)
            self.assertEqual(sheet.cell(row=2, column=2).value, "Sea of Stone")
            self.assertEqual(sheet.cell(row=2, column=4).value, 51.77)

            summary = book["Сводка"]
            values = dict(summary.iter_rows(min_row=1, values_only=True))
            self.assertEqual(values["Всего записей"], 3)
            self.assertEqual(values["Записей с ценой"], 3)
            self.assertEqual(values["Минимальная цена"], 12.99)
            self.assertEqual(values["Максимальная цена"], 51.77)
            self.assertEqual(values["Средняя цена"], 28.09)
        finally:
            out_path.unlink(missing_ok=True)


if __name__ == "__main__":
    unittest.main()

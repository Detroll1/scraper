"""Ядро демо-парсера: загрузка источника, разбор записей, вывод в Excel и JSON.

Только requests + BeautifulSoup, без headless-браузеров. Excel — через openpyxl.
"""

from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = Path(__file__).resolve().parent
DEMO_HTML = BASE_DIR / "data" / "demo_catalog.html"

DEFAULT_TIMEOUT = 20
DEFAULT_USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/124.0.0.0 Safari/537.36"
)

# Пресеты источников: имя → (тип, адрес, подпись для отчёта).
PRESETS = {
    "demo": ("file", DEMO_HTML, "Демо-каталог книг (локальный офлайн-источник, data/demo_catalog.html)"),
    "books": ("url", "https://books.toscrape.com/", "Books to Scrape — открытый каталог книг (books.toscrape.com)"),
}

# Цены вида «£51.77», «$12», «499,00 ₽», «299 руб.», «19,99 USD».
PRICE_RE = re.compile(
    r"(?:[$£€₽]\s?(\d+(?:[.,]\d{1,2})?)"
    r"|(\d+(?:[.,]\d{1,2})?)\s?[$£€₽]"
    r"|\b(\d+(?:[.,]\d{1,2})?)\s?(?:руб\.?|USD|RUB|EUR)\b)",
    re.IGNORECASE,
)

DATA_HEADERS = ("№", "Наименование", "Ссылка", "Цена", "Мета")

EXIT_SOURCE_ERROR = 2


class ScrapingError(Exception):
    """Понятная пользователю ошибка источника (сеть, 404, файл).

    CLI печатает её без трейсбека и завершается с кодом exit_code != 0.
    """

    def __init__(self, message: str, exit_code: int = EXIT_SOURCE_ERROR):
        super().__init__(message)
        self.exit_code = exit_code


@dataclass
class Record:
    """Одна запись каталога: наименование, ссылка, цена и произвольная мета."""

    title: str
    link: str
    price: float | None
    meta: str | None

    def to_dict(self) -> dict:
        return {"title": self.title, "link": self.link, "price": self.price, "meta": self.meta}


def _settings() -> dict:
    """Читает config.py (копия config.example.py), если он лежит рядом."""
    try:
        import config  # noqa: PLC0415
    except ImportError:
        return {}
    keys = ("TIMEOUT", "USER_AGENT")
    return {key: getattr(config, key) for key in keys if hasattr(config, key)}


def parse_price(text: str | None) -> float | None:
    """Достаёт число из строки с ценой: «£51.77» → 51.77, «499,00 ₽» → 499.0."""
    if not text:
        return None
    match = PRICE_RE.search(text)
    if match is None:
        return None
    number = next(group for group in match.groups() if group)
    try:
        return float(number.replace(",", "."))
    except ValueError:
        return None


def load_source(source: str, timeout: int | None = None) -> tuple[str, str, str]:
    """Загружает источник и возвращает (html, base_url, подпись источника).

    source — имя пресета (demo, books), URL http(s):// или путь к HTML-файлу.
    Любая проблема доступа превращается в ScrapingError с понятным текстом.
    """
    settings = _settings()
    timeout = timeout or settings.get("TIMEOUT") or DEFAULT_TIMEOUT
    user_agent = settings.get("USER_AGENT") or DEFAULT_USER_AGENT

    preset = PRESETS.get(source.strip().lower())
    if preset is not None:
        kind, address, label = preset
        if kind == "file":
            address = Path(address)
            return _read_local_file(address), _dir_uri(address.parent), label
        return _fetch_url(address, timeout=timeout, user_agent=user_agent), address, label

    if re.match(r"^https?://", source, re.IGNORECASE):
        return _fetch_url(source, timeout=timeout, user_agent=user_agent), source, source

    path = Path(source)
    if path.is_file() and path.suffix.lower() in {".html", ".htm"}:
        return _read_local_file(path), _dir_uri(path.resolve().parent), f"Локальный файл {path}"

    raise ScrapingError(
        f"Не понял источник «{source}»: ожидаю URL http(s)://, "
        f"пресет ({', '.join(PRESETS)}) или путь к HTML-файлу."
    )


def _dir_uri(directory: Path) -> str:
    """file://-URI папки со слэшем на конце — чтобы urljoin собирал абсолютные ссылки."""
    return directory.as_uri().rstrip("/") + "/"


def _read_local_file(path: Path) -> str:
    if not path.is_file():
        raise ScrapingError(f"Файл источника не найден: {path}")
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError as exc:
        raise ScrapingError(f"Файл {path} не в UTF-8, парсер его не читает") from exc
    except OSError as exc:
        raise ScrapingError(f"Не удалось прочитать файл {path}: {exc.strerror or exc}") from exc


def _fetch_url(url: str, timeout: int, user_agent: str) -> str:
    try:
        response = requests.get(
            url,
            timeout=timeout,
            headers={"User-Agent": user_agent, "Accept-Language": "ru,en;q=0.8"},
        )
    except requests.exceptions.Timeout as exc:
        raise ScrapingError(f"Источник не ответил за {timeout} сек (таймаут): {url}") from exc
    except requests.exceptions.ConnectionError as exc:
        host = re.sub(r"^https?://([^/]+).*$", r"\1", url, flags=re.IGNORECASE)
        raise ScrapingError(f"Нет соединения с {host} — проверьте сеть и адрес: {url}") from exc
    except requests.exceptions.RequestException as exc:
        raise ScrapingError(
            f"Не удалось загрузить страницу {url} ({exc.__class__.__name__})"
        ) from exc

    if response.status_code == 404:
        raise ScrapingError(f"Источник вернул 404 (страница не найдена): {url}")
    if response.status_code != 200:
        raise ScrapingError(f"Источник вернул HTTP {response.status_code}: {url}")

    if not response.encoding or response.encoding.lower() == "iso-8859-1":
        response.encoding = response.apparent_encoding
    return response.text


def parse_records(html: str, base_url: str) -> list[Record]:
    """Разбирает HTML в список записей: сперва карточки каталога, потом общий разбор.

    Повторяющиеся записи (одинаковые наименование и ссылка) отбрасываются.
    """
    soup = BeautifulSoup(html, "html.parser")
    records = _catalog_records(soup, base_url)
    if not records:
        records = _generic_records(soup, base_url)

    unique: list[Record] = []
    seen: set[tuple[str, str]] = set()
    for record in records:
        key = (record.title, record.link)
        if key in seen:
            continue
        seen.add(key)
        unique.append(record)
    return unique


def _catalog_records(soup: BeautifulSoup, base_url: str) -> list[Record]:
    """Карточки вида article.product_pod (books.toscrape.com и наш демо-каталог)."""
    records: list[Record] = []
    for card in soup.select("article.product_pod"):
        anchor = card.select_one("h3 a")
        if anchor is None:
            continue
        title = (anchor.get("title") or anchor.get_text(" ", strip=True) or "").strip()
        if len(title) < 3:
            continue
        link = urljoin(base_url, anchor.get("href", ""))

        price_tag = card.select_one(".price_color")
        price = parse_price(price_tag.get_text() if price_tag else "")

        meta_parts: list[str] = []
        rating = card.select_one("p.star-rating")
        if rating is not None:
            stars = [cls for cls in rating.get("class", []) if cls != "star-rating"]
            if stars:
                meta_parts.append(f"Рейтинг: {stars[0]}")
        availability = card.select_one(".availability")
        if availability is not None:
            meta_parts.append(availability.get_text(" ", strip=True))

        records.append(Record(title=title, link=link, price=price, meta="; ".join(meta_parts) or None))
    return records


def _generic_records(soup: BeautifulSoup, base_url: str) -> list[Record]:
    """Общий разбор без карточек: заголовки и списки со ссылками (например, топы, оглавления)."""
    records: list[Record] = []
    seen: set[tuple[str, str]] = set()
    for anchor in soup.select("h2 a[href], h3 a[href], h4 a[href], ol li a[href]"):
        title = anchor.get_text(" ", strip=True)
        if len(title) < 3:
            continue
        link = urljoin(base_url, anchor.get("href", ""))
        key = (title, link)
        if key in seen:
            continue
        seen.add(key)

        container = anchor.find_parent(["article", "li", "div"])
        card_text = container.get_text(" ", strip=True) if container is not None else ""
        price = parse_price(card_text)
        meta = card_text[:180] if card_text and card_text != title else None
        records.append(Record(title=title, link=link, price=price, meta=meta))
    return records


def _summary_rows(records: list[Record], source_label: str) -> list[tuple[str, object]]:
    prices = [record.price for record in records if record.price is not None]
    titles = {record.title for record in records}
    rows: list[tuple[str, object]] = [
        ("Источник", source_label),
        ("Дата генерации", datetime.now().strftime("%Y-%m-%d %H:%M")),
        ("Всего записей", len(records)),
        ("Уникальных наименований", len(titles)),
        ("Записей с ценой", len(prices)),
    ]
    if prices:
        rows += [
            ("Минимальная цена", min(prices)),
            ("Максимальная цена", max(prices)),
            ("Средняя цена", round(sum(prices) / len(prices), 2)),
        ]
    else:
        rows += [
            ("Минимальная цена", "—"),
            ("Максимальная цена", "—"),
            ("Средняя цена", "—"),
        ]
    return rows


def build_workbook(records: list[Record], out_path: Path | str, source_label: str) -> None:
    """Пишет xlsx: лист «Данные» (записи) + лист «Сводка» (агрегаты)."""
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Данные"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="4472C4")
    for column, title in enumerate(DATA_HEADERS, start=1):
        cell = data_sheet.cell(row=1, column=column, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for index, record in enumerate(records, start=1):
        row = index + 1
        data_sheet.cell(row=row, column=1, value=index)
        data_sheet.cell(row=row, column=2, value=record.title)
        link_cell = data_sheet.cell(row=row, column=3, value=record.link)
        if record.link:
            link_cell.hyperlink = record.link
            link_cell.font = Font(color="0563C1", underline="single")
        if record.price is not None:
            price_cell = data_sheet.cell(row=row, column=4, value=record.price)
            price_cell.number_format = "0.00"
        data_sheet.cell(row=row, column=5, value=record.meta)

    for column, width in enumerate((6, 48, 56, 12, 40), start=1):
        data_sheet.column_dimensions[get_column_letter(column)].width = width
    data_sheet.freeze_panes = "A2"

    summary_sheet = workbook.create_sheet("Сводка")
    for row, (label, value) in enumerate(_summary_rows(records, source_label), start=1):
        label_cell = summary_sheet.cell(row=row, column=1, value=label)
        label_cell.font = Font(bold=True)
        value_cell = summary_sheet.cell(row=row, column=2, value=value)
        if isinstance(value, float):
            value_cell.number_format = "0.00"
    summary_sheet.column_dimensions["A"].width = 26
    summary_sheet.column_dimensions["B"].width = 64

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def write_json(records: list[Record], json_path: Path | str, source_label: str) -> None:
    """Дополнительный вывод в JSON (--json FILE)."""
    payload = {
        "source": source_label,
        "generated": datetime.now().isoformat(timespec="seconds"),
        "count": len(records),
        "records": [record.to_dict() for record in records],
    }
    json_path = Path(json_path)
    json_path.parent.mkdir(parents=True, exist_ok=True)
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
